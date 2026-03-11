from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook
from models import ImportSummary


class ExcelProcessingError(Exception):
    pass


def process_mtr_excel(file_path: Path) -> ImportSummary:
    try:
        wb = load_workbook(filename=file_path, data_only=True)
    except Exception as exc:
        raise ExcelProcessingError(f"Não foi possível abrir o Excel: {exc}") from exc

    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ExcelProcessingError("A planilha está vazia.")

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data_rows = [dict(zip(headers, row)) for row in rows[1:]]

    total_rows = len(data_rows)
    saved_rows = 0
    missing_mtr_code_rows = 0
    mtr_codes: set[str] = set()

    for row in data_rows:
        status = str(row.get("Situação") or "").strip().upper()
        if status != "SALVO":
            continue

        saved_rows += 1
        mtr_code = str(row.get("Nº MTR") or "").strip()
        if not mtr_code:
            missing_mtr_code_rows += 1
            continue
        mtr_codes.add(mtr_code)

    return ImportSummary(
        total_rows=total_rows,
        saved_rows=saved_rows,
        unique_mtrs=len(mtr_codes),
        missing_mtr_code_rows=missing_mtr_code_rows,
    )
